import unittest
import tempfile
from pathlib import Path
from PIL import Image, ImageDraw
import docx
import pypdf

from shiftdrop.core.dispatcher import execute_action
from shiftdrop.core.ffmpeg_helper import run_ffmpeg

class TestAllActionsExhaustive(unittest.TestCase):
    """
    Exhaustive verification: Tests 1-by-1 every single conversion action supported in ShiftDrop.
    Verifies that the generated file exists, is non-empty, and is in the same directory.
    """
    @classmethod
    def setUpClass(cls):
        cls.temp_dir = tempfile.TemporaryDirectory()
        cls.dir = Path(cls.temp_dir.name)

        # 1. Image
        cls.img_png = cls.dir / "test_input_image.png"
        img = Image.new("RGBA", (100, 100), (255, 100, 0, 255))
        img.save(cls.img_png)

        cls.img_jpg = cls.dir / "test_input_image.jpg"
        img.convert("RGB").save(cls.img_jpg)

        # 2. Audio
        cls.audio_wav = cls.dir / "test_input_audio.wav"
        run_ffmpeg(["-f", "lavfi", "-i", "sine=frequency=440:duration=1.0", "-c:a", "pcm_s16le", str(cls.audio_wav)])

        # 3. Video
        cls.video_mp4 = cls.dir / "test_input_video.mp4"
        run_ffmpeg([
            "-f", "lavfi", "-i", "testsrc=duration=1:size=160x120:rate=10",
            "-f", "lavfi", "-i", "sine=frequency=440:duration=1",
            "-c:v", "libx264", "-pix_fmt", "yuv420p", "-c:a", "aac",
            str(cls.video_mp4)
        ])

        # 4. PDF (with text)
        cls.pdf_doc = cls.dir / "test_input_doc.pdf"
        p_img = Image.new("RGB", (400, 600), (255, 255, 255))
        draw = ImageDraw.Draw(p_img)
        draw.text((40, 40), "ShiftDrop Exhaustive Test PDF", fill=(0, 0, 0))
        p_img.save(cls.pdf_doc, "PDF")

        # 5. Word Document (.docx)
        cls.docx_doc = cls.dir / "test_input_word.docx"
        d = docx.Document()
        d.add_heading("ShiftDrop Exhaustive Test Word", 0)
        d.add_paragraph("Párrafo de prueba para conversión exhaustiva.")
        d.save(str(cls.docx_doc))

        # 6. Plain Text (.txt)
        cls.txt_doc = cls.dir / "test_input_notes.txt"
        with open(cls.txt_doc, "w", encoding="utf-8") as f:
            f.write("Línea 1: Contenido de prueba en texto plano.\nLínea 2: Segunda línea.")

    @classmethod
    def tearDownClass(cls):
        cls.temp_dir.cleanup()

    def _verify_result(self, results: list[Path], expected_ext: str = None):
        self.assertTrue(len(results) > 0, "No se generaron archivos de salida.")
        for r in results:
            self.assertTrue(r.exists(), f"El archivo generado {r} no existe.")
            self.assertGreater(r.stat().st_size, 0, f"El archivo generado {r} está vacío.")
            # Verify file is located in or inside the directory of the original file
            self.assertTrue(
                str(r.resolve()).lower().startswith(str(self.dir.resolve()).lower()),
                f"{r} no está en el directorio original."
            )
            if expected_ext:
                self.assertEqual(r.suffix.lower(), expected_ext.lower(), f"Extensión incorrecta para {r}")

    # ==================== IMAGE ACTIONS ====================
    def test_01_img_to_webp(self):
        res = execute_action("img_to_webp", [self.img_png])
        self._verify_result(res, ".webp")

    def test_02_img_to_png(self):
        res = execute_action("img_to_png", [self.img_jpg])
        self._verify_result(res, ".png")

    def test_03_img_to_jpg(self):
        res = execute_action("img_to_jpg", [self.img_png])
        self._verify_result(res, ".jpg")

    def test_04_img_to_ico(self):
        res = execute_action("img_to_ico", [self.img_png])
        self._verify_result(res, ".ico")

    def test_05_img_to_bmp(self):
        res = execute_action("img_to_bmp", [self.img_png])
        self._verify_result(res, ".bmp")

    def test_06_img_to_tiff(self):
        res = execute_action("img_to_tiff", [self.img_png])
        self._verify_result(res, ".tiff")

    def test_07_img_to_gif(self):
        res = execute_action("img_to_gif", [self.img_png])
        self._verify_result(res, ".gif")

    def test_08_img_to_pdf_single(self):
        res = execute_action("img_to_pdf", [self.img_png])
        self._verify_result(res, ".pdf")

    def test_09_img_to_pdf_multi(self):
        res = execute_action("img_to_pdf", [self.img_png, self.img_jpg])
        self._verify_result(res, ".pdf")

    def test_10_img_compress(self):
        res = execute_action("img_compress", [self.img_png], params={"quality": 60})
        self._verify_result(res)

    def test_11_img_crop_resize_aspect(self):
        res = execute_action("img_crop_resize", [self.img_png], params={"aspect_ratio": "1:1"})
        self._verify_result(res)

    def test_12_img_crop_resize_scale(self):
        res = execute_action("img_crop_resize", [self.img_png], params={"scale_percent": 0.5})
        self._verify_result(res)

    # ==================== VIDEO ACTIONS ====================
    def test_13_vid_to_mp4(self):
        res = execute_action("vid_to_mp4", [self.video_mp4])
        self._verify_result(res, ".mp4")

    def test_14_vid_to_webm(self):
        res = execute_action("vid_to_webm", [self.video_mp4])
        self._verify_result(res, ".webm")

    def test_15_vid_to_mkv(self):
        res = execute_action("vid_to_mkv", [self.video_mp4])
        self._verify_result(res, ".mkv")

    def test_16_vid_to_mov(self):
        res = execute_action("vid_to_mov", [self.video_mp4])
        self._verify_result(res, ".mov")

    def test_17_vid_to_avi(self):
        res = execute_action("vid_to_avi", [self.video_mp4])
        self._verify_result(res, ".avi")

    def test_18_vid_to_gif(self):
        res = execute_action("vid_to_gif", [self.video_mp4], params={"fps": 10, "width": 120})
        self._verify_result(res, ".gif")

    def test_19_vid_extract_mp3(self):
        res = execute_action("vid_extract_mp3", [self.video_mp4])
        self._verify_result(res, ".mp3")

    def test_20_vid_extract_wav(self):
        res = execute_action("vid_extract_wav", [self.video_mp4])
        self._verify_result(res, ".wav")

    def test_21_vid_compress(self):
        res = execute_action("vid_compress", [self.video_mp4], params={"crf": 30})
        self._verify_result(res)

    def test_22_vid_trim(self):
        res = execute_action("vid_trim", [self.video_mp4], params={"start_sec": 0, "duration_sec": 0.5})
        self._verify_result(res)

    def test_23_vid_mute(self):
        res = execute_action("vid_mute", [self.video_mp4])
        self._verify_result(res)

    # ==================== AUDIO ACTIONS ====================
    def test_24_aud_to_mp3(self):
        res = execute_action("aud_to_mp3", [self.audio_wav])
        self._verify_result(res, ".mp3")

    def test_25_aud_to_wav(self):
        res = execute_action("aud_to_wav", [self.audio_wav])
        self._verify_result(res, ".wav")

    def test_26_aud_to_flac(self):
        res = execute_action("aud_to_flac", [self.audio_wav])
        self._verify_result(res, ".flac")

    def test_27_aud_to_m4a(self):
        res = execute_action("aud_to_m4a", [self.audio_wav])
        self._verify_result(res, ".m4a")

    def test_28_aud_to_aac(self):
        res = execute_action("aud_to_aac", [self.audio_wav])
        self._verify_result(res, ".m4a")

    def test_29_aud_to_ogg(self):
        res = execute_action("aud_to_ogg", [self.audio_wav])
        self._verify_result(res, ".ogg")

    def test_30_aud_to_opus(self):
        res = execute_action("aud_to_opus", [self.audio_wav])
        self._verify_result(res, ".opus")

    def test_31_aud_normalize(self):
        res = execute_action("aud_normalize", [self.audio_wav])
        self._verify_result(res)

    def test_32_aud_trim(self):
        res = execute_action("aud_trim", [self.audio_wav], params={"start_sec": 0, "duration_sec": 0.5})
        self._verify_result(res)

    # ==================== PDF ACTIONS ====================
    def test_33_pdf_merge(self):
        res = execute_action("pdf_merge", [self.pdf_doc, self.pdf_doc])
        self._verify_result(res, ".pdf")

    def test_34_pdf_to_images_png(self):
        res = execute_action("pdf_to_images", [self.pdf_doc], params={"format": "png"})
        self._verify_result(res, ".png")

    def test_35_pdf_to_images_jpg(self):
        res = execute_action("pdf_to_images", [self.pdf_doc], params={"format": "jpg"})
        self._verify_result(res, ".jpg")

    def test_36_pdf_compress(self):
        res = execute_action("pdf_compress", [self.pdf_doc])
        self._verify_result(res, ".pdf")

    def test_37_pdf_split(self):
        res = execute_action("pdf_split", [self.pdf_doc], params={"pages": 1})
        self.assertTrue(len(res) > 0)
        for r in res:
            self.assertTrue(r.exists())
            self.assertGreater(r.stat().st_size, 0)

    def test_38_pdf_to_docx(self):
        res = execute_action("pdf_to_docx", [self.pdf_doc])
        self._verify_result(res, ".docx")

    # ==================== DOCUMENT ACTIONS ====================
    def test_39_doc_to_pdf_from_docx(self):
        res = execute_action("doc_to_pdf", [self.docx_doc])
        self._verify_result(res, ".pdf")

    def test_40_doc_to_pdf_from_txt(self):
        res = execute_action("doc_to_pdf", [self.txt_doc])
        self._verify_result(res, ".pdf")

    def test_41_doc_to_txt(self):
        res = execute_action("doc_to_txt", [self.docx_doc])
        self._verify_result(res, ".txt")

    def test_42_doc_to_images(self):
        res = execute_action("doc_to_images", [self.docx_doc])
        self._verify_result(res, ".png")

    # ==================== COMPRESS FALLBACKS ====================
    def test_43_compress_image(self):
        res = execute_action("compress", [self.img_png])
        self._verify_result(res)

    def test_44_compress_video(self):
        res = execute_action("compress", [self.video_mp4])
        self._verify_result(res)

    def test_45_compress_pdf(self):
        res = execute_action("compress", [self.pdf_doc])
        self._verify_result(res)

    def test_46_compress_docs_to_zip(self):
        res = execute_action("compress", [self.docx_doc, self.txt_doc])
        self._verify_result(res, ".zip")

if __name__ == "__main__":
    unittest.main()
