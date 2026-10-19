import unittest
import tempfile
from pathlib import Path
import docx
import pypdf

from shiftdrop.core.doc_ops import doc_to_pdf, pdf_to_docx, doc_to_txt
from shiftdrop.core.pdf_ops import _validate_pdf_file
from shiftdrop.core.dispatcher import detect_file_category, execute_action
from shiftdrop.ui.slice_definition import get_slices_for_category, DOCUMENT_SLICES, DOC_MULTI_SLICES

class TestDocOps(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.temp_path = Path(self.temp_dir.name)

        # Create a sample .docx file (which starts with PK\x03\x04)
        self.docx_file = self.temp_path / "sample.docx"
        d = docx.Document()
        d.add_heading("ShiftDrop Document Test", 0)
        d.add_paragraph("Este es un párrafo de prueba para conversión.")
        d.save(str(self.docx_file))

        # Create a sample .txt file
        self.txt_file = self.temp_path / "notes.txt"
        with open(self.txt_file, "w", encoding="utf-8") as f:
            f.write("Línea 1: Documento de texto.\nLínea 2: Segunda línea de prueba.")

        # Create a valid PDF file
        self.pdf_file = self.temp_path / "sample.pdf"
        writer = pypdf.PdfWriter()
        writer.add_blank_page(width=100, height=100)
        with open(self.pdf_file, "wb") as f:
            writer.write(f)

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_detect_category_documents(self):
        self.assertEqual(detect_file_category([self.docx_file]), "document")
        self.assertEqual(detect_file_category([self.txt_file]), "document")
        self.assertEqual(detect_file_category([self.docx_file, self.txt_file]), "doc_multi")
        self.assertEqual(detect_file_category([self.pdf_file]), "pdf")

    def test_slices_for_document_category(self):
        single_slices = get_slices_for_category("document")
        self.assertEqual(single_slices, DOCUMENT_SLICES)
        action_ids = [s.action_id for s in single_slices if s.action_id]
        self.assertIn("doc_to_pdf", action_ids)
        self.assertIn("doc_to_txt", action_ids)
        self.assertIn("doc_to_images", action_ids)
        self.assertIn("compress", action_ids)

        multi_slices = get_slices_for_category("doc_multi")
        self.assertEqual(multi_slices, DOC_MULTI_SLICES)
        multi_action_ids = [s.action_id for s in multi_slices if s.action_id]
        self.assertIn("doc_to_pdf", multi_action_ids)
        self.assertIn("doc_to_txt", multi_action_ids)

    def test_validate_pdf_rejects_pk_docx_header(self):
        # Must catch PK\x03\x04 zip header from docx and raise informative ValueError
        with self.assertRaises(ValueError) as ctx:
            _validate_pdf_file(self.docx_file)
        self.assertIn("no un PDF", str(ctx.exception))

    def test_doc_to_txt_from_docx(self):
        out_txt = doc_to_txt(self.docx_file)
        self.assertTrue(out_txt.exists())
        self.assertEqual(out_txt.parent, self.temp_path)
        with open(out_txt, "r", encoding="utf-8") as f:
            content = f.read()
        self.assertIn("ShiftDrop Document Test", content)
        self.assertIn("Este es un párrafo de prueba", content)

    def test_doc_to_pdf_from_txt(self):
        out_pdf = doc_to_pdf(self.txt_file)
        self.assertTrue(out_pdf.exists())
        self.assertEqual(out_pdf.parent, self.temp_path)
        # Should be a valid PDF
        _validate_pdf_file(out_pdf)

    def test_doc_to_pdf_from_docx(self):
        # Native Word COM or fallback
        try:
            out_pdf = doc_to_pdf(self.docx_file)
            self.assertTrue(out_pdf.exists())
            self.assertEqual(out_pdf.parent, self.temp_path)
            _validate_pdf_file(out_pdf)
        except RuntimeError:
            pass  # If environment lacks MS Word COM

    def test_pdf_to_docx(self):
        out_docx = pdf_to_docx(self.pdf_file)
        self.assertTrue(out_docx.exists())
        self.assertEqual(out_docx.parent, self.temp_path)
        doc = docx.Document(str(out_docx))
        self.assertTrue(len(doc.paragraphs) > 0)

    def test_dispatcher_document_actions(self):
        # Dispatcher doc_to_txt
        res = execute_action("doc_to_txt", [self.docx_file])
        self.assertEqual(len(res), 1)
        self.assertTrue(res[0].exists())
        self.assertEqual(res[0].suffix, ".txt")
        self.assertEqual(res[0].parent, self.temp_path)

        # Dispatcher compress on documents creates a zip in the same directory
        res_zip = execute_action("compress", [self.docx_file, self.txt_file])
        self.assertEqual(len(res_zip), 1)
        self.assertTrue(res_zip[0].exists())
        self.assertEqual(res_zip[0].suffix, ".zip")
        self.assertEqual(res_zip[0].parent, self.temp_path)

if __name__ == "__main__":
    unittest.main()
