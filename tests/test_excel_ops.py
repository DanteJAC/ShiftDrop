"""
Unit tests for Excel and spreadsheet conversion operations.
"""
import unittest
import tempfile
import json
import csv
from pathlib import Path
import openpyxl
import docx

from shiftdrop.core.excel_ops import (
    excel_to_pdf, excel_to_docx, excel_to_csv,
    excel_to_json, excel_to_html, csv_to_excel,
    read_spreadsheet_data
)
from shiftdrop.core.dispatcher import (
    inspect_file_type, detect_file_category, execute_action
)


class TestExcelOps(unittest.TestCase):
    def setUp(self):
        self.tmp_dir = tempfile.TemporaryDirectory()
        self.dir_path = Path(self.tmp_dir.name)

        # Create sample XLSX
        self.sample_xlsx = self.dir_path / "presupuesto.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Ventas 2026"
        ws.append(["Producto", "Cantidad", "Precio Unitario", "Total"])
        ws.append(["Laptop Dell", 5, 850.00, 4250.00])
        ws.append(["Monitor 4K", 10, 320.50, 3205.00])
        ws.append(["Teclado Mecánico", 25, 65.00, 1625.00])
        wb.save(str(self.sample_xlsx))

        # Create sample CSV
        self.sample_csv = self.dir_path / "usuarios.csv"
        with open(self.sample_csv, "w", encoding="utf-8", newline="") as f:
            writer = csv.writer(f)
            writer.writerow(["ID", "Nombre", "Departamento", "Activo"])
            writer.writerow(["101", "Dante", "Ingeniería", "True"])
            writer.writerow(["102", "Elena", "Diseño", "True"])
            writer.writerow(["103", "Carlos", "Marketing", "False"])

    def tearDown(self):
        self.tmp_dir.cleanup()

    def test_inspect_file_type_spreadsheet(self):
        self.assertEqual(inspect_file_type(self.sample_xlsx), "spreadsheet")
        self.assertEqual(inspect_file_type(self.sample_csv), "spreadsheet")

    def test_detect_file_category_multi(self):
        cat_single = detect_file_category([self.sample_xlsx])
        self.assertEqual(cat_single, "spreadsheet")

        cat_multi = detect_file_category([self.sample_xlsx, self.sample_csv])
        self.assertEqual(cat_multi, "spreadsheet_multi")

    def test_read_spreadsheet_data(self):
        title, rows = read_spreadsheet_data(self.sample_xlsx)
        self.assertEqual(title, "Ventas 2026")
        self.assertEqual(len(rows), 4)
        self.assertEqual(rows[0][0], "Producto")
        self.assertEqual(rows[1][0], "Laptop Dell")

    def test_excel_to_pdf(self):
        out_pdf = excel_to_pdf(self.sample_xlsx)
        self.assertTrue(out_pdf.exists())
        self.assertGreater(out_pdf.stat().st_size, 500)
        with open(out_pdf, "rb") as f:
            header = f.read(5)
            self.assertEqual(header, b"%PDF-")

    def test_excel_to_docx(self):
        out_docx = excel_to_docx(self.sample_xlsx)
        self.assertTrue(out_docx.exists())
        self.assertGreater(out_docx.stat().st_size, 1000)
        doc = docx.Document(str(out_docx))
        self.assertGreater(len(doc.tables), 0)
        table = doc.tables[0]
        self.assertEqual(table.rows[0].cells[0].text, "Producto")
        self.assertEqual(table.rows[1].cells[0].text, "Laptop Dell")

    def test_excel_to_csv(self):
        out_csv = excel_to_csv(self.sample_xlsx)
        self.assertTrue(out_csv.exists())
        with open(out_csv, "r", encoding="utf-8-sig") as f:
            reader = list(csv.reader(f))
            self.assertEqual(reader[0][0], "Producto")
            self.assertEqual(reader[1][0], "Laptop Dell")

    def test_excel_to_json(self):
        out_json = excel_to_json(self.sample_xlsx)
        self.assertTrue(out_json.exists())
        data = json.loads(out_json.read_text(encoding="utf-8"))
        self.assertEqual(len(data), 3)
        self.assertEqual(data[0]["Producto"], "Laptop Dell")
        self.assertEqual(str(data[0]["Cantidad"]), "5")

    def test_excel_to_html(self):
        out_html = excel_to_html(self.sample_xlsx)
        self.assertTrue(out_html.exists())
        content = out_html.read_text(encoding="utf-8")
        self.assertIn("<table", content)
        self.assertIn("Laptop Dell", content)
        self.assertIn("Ventas 2026", content)

    def test_csv_to_excel(self):
        out_xlsx = csv_to_excel(self.sample_csv)
        self.assertTrue(out_xlsx.exists())
        wb = openpyxl.load_workbook(out_xlsx)
        ws = wb.active
        self.assertEqual(ws.cell(1, 1).value, "ID")
        self.assertEqual(ws.cell(2, 2).value, "Dante")
        wb.close()

    def test_dispatcher_excel_actions(self):
        # 1. Dispatcher excel_to_pdf
        res_pdf = execute_action("excel_to_pdf", [self.sample_xlsx])
        self.assertEqual(len(res_pdf), 1)
        self.assertTrue(res_pdf[0].exists())

        # 2. Dispatcher excel_to_docx
        res_docx = execute_action("excel_to_docx", [self.sample_xlsx])
        self.assertEqual(len(res_docx), 1)
        self.assertTrue(res_docx[0].exists())

        # 3. Dispatcher excel_to_csv
        res_csv = execute_action("excel_to_csv", [self.sample_xlsx])
        self.assertEqual(len(res_csv), 1)
        self.assertTrue(res_csv[0].exists())

        # 4. Dispatcher excel_to_json
        res_json = execute_action("excel_to_json", [self.sample_xlsx])
        self.assertEqual(len(res_json), 1)
        self.assertTrue(res_json[0].exists())

        # 5. Dispatcher excel_to_html
        res_html = execute_action("excel_to_html", [self.sample_xlsx])
        self.assertEqual(len(res_html), 1)
        self.assertTrue(res_html[0].exists())

        # 6. Dispatcher csv_to_excel
        res_xlsx = execute_action("csv_to_excel", [self.sample_csv])
        self.assertEqual(len(res_xlsx), 1)
        self.assertTrue(res_xlsx[0].exists())


if __name__ == "__main__":
    unittest.main()
